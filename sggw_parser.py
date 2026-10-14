"""
Parser siatki planów zajęć SGGW z arkuszy Excel (.xlsx).
"""
import re
import os
import logging
from datetime import datetime, timedelta
from typing import Dict, List, Optional, Tuple, Any, Union
import openpyxl

logger = logging.getLogger("sggw_parser")

DAY_NAME_MAP = {
    "poniedziałek": "PON",
    "poniedzialek": "PON",
    "poniedz.": "PON",
    "poniedz": "PON",
    "pon": "PON",
    "wtorek": "WT",
    "wt": "WT",
    "środa": "ŚR",
    "sroda": "ŚR",
    "śr": "ŚR",
    "sr": "ŚR",
    "czwartek": "CZW",
    "czw": "CZW",
    "piątek": "PT",
    "piatek": "PT",
    "pt": "PT"
}


def parse_time_bounds(text: str, default_start: str, default_end: str) -> Tuple[str, str]:
    """
    Wyznacza godziny rozpoczęcia i zakończenia zajęć.
    Jeśli w tekście podano jawne 'od XX:YY do AA:BB' lub 'od XX:YY' / 'do AA:BB', ma to priorytet.
    """
    start_time = default_start
    end_time = default_end

    # Format: od HH:MM do HH:MM lub HH:MM - HH:MM
    range_match = re.search(r'(?:od\s*)?(\d{1,2}[:.]\d{2})\s*(?:do|-)\s*(\d{1,2}[:.]\d{2})', text, re.IGNORECASE)
    if range_match:
        s = range_match.group(1).replace(".", ":")
        e = range_match.group(2).replace(".", ":")
        return (f"{int(s.split(':')[0]):02d}:{s.split(':')[1]}", f"{int(e.split(':')[0]):02d}:{e.split(':')[1]}")

    # Format: tylko "od HH:MM"
    from_match = re.search(r'\bod\s*(\d{1,2}[:.]\d{2})', text, re.IGNORECASE)
    if from_match:
        s = from_match.group(1).replace(".", ":")
        start_time = f"{int(s.split(':')[0]):02d}:{s.split(':')[1]}"

    # Format: tylko "do HH:MM"
    to_match = re.search(r'\bdo\s*(\d{1,2}[:.]\d{2})', text, re.IGNORECASE)
    if to_match:
        e = to_match.group(1).replace(".", ":")
        end_time = f"{int(e.split(':')[0]):02d}:{e.split(':')[1]}"

    return (start_time, end_time)


def extract_groups_from_text(text: str, all_groups: List[str]) -> List[str]:
    """
    Identyfikuje grupy ćwiczeniowe z tekstu komórki.
    Jeśli brak wzmianki o grupie (lub wykład dla rocznika/fakultet), zwraca wszystkie grupy.
    """
    clean_text = text.replace("\n", " ")
    
    # Wykrywanie wzorców grup: 'gr. 1, 2', 'Gr 3,4,7', 'gr. 5, 6', 'gr 1'
    group_matches = re.findall(r'(?:gr\.?|grupa|grup)\s*([0-9\s,]+)', clean_text, re.IGNORECASE)
    if group_matches:
        target_group_nums = set()
        for gm in group_matches:
            # Rozbicie po przecinkach i spacjach
            nums = re.findall(r'\b([1-9])\b', gm)
            target_group_nums.update(int(n) for n in nums)
        
        matched_groups = [f"gr. {n}" for n in sorted(target_group_nums) if f"gr. {n}" in all_groups]
        if matched_groups:
            return matched_groups

    # Jeśli nie wykryto żadnej specyficznej grupy -> zajęcia wspólne / dla wszystkich
    return list(all_groups)


def classify_sggw_form(text: str) -> str:
    """Klasyfikuje formę zajęć na podstawie treści komórki."""
    lower = text.lower()
    if "fakultet" in lower:
        return "fakultet"
    if "ćw" in lower or "ćwiczen" in lower:
        return "cwiczenia"
    if "wykład" in lower or "wyklad" in lower or lower.strip().endswith(" w.") or " w." in lower:
        return "wyklad"
    if "lab" in lower:
        return "laboratorium"
    return "cwiczenia"


ROOM_PATTERNS = [
    r'(\b\d{2,4}\s*/\s*Aula\s*\d*(?:\s*b\.?\s*\d+)?)',
    r'(Wolica\s*sala\s*\d+,\s*aula\s*\d+)',
    r'(Wolica\s*sala\s*\d+)',
    r'(sala\s*\d+,\s*aula\s*\d+)',
    r'(sala\s*109/ambulatorium/Obory)',
    r'(sala109/ambulatorium/Obory)',
    r'(06/ambul(?:\s*b\.?\s*\d+)?)',
    r'(sale\s*zakładu(?:\s*b\.?\s*\d+)?)',
    r'(sala\s*sekcyjna(?:\s*/\s*sala\s*ćwiczeniowa)?)',
    r'(aud\s*\d*(?:\s*b\.?\s*\d+)?)',
    r'(aud\s*b\.?\s*\d+)',
    r'(Aula\s*\d*(?:\s*b\.?\s*\d+)?)',
    r'(Audytorium(?:\s*b\.?\s*\d+)?)',
    r'((?:b\.?\s*\d+\s+)?sala\s*\d+(?:[a-z])?(?:/\d+)?(?:\s*b\.?\s*\d+)?)',
    r'(s\s*\d+\s*b\.?\s*\d+)',
    r'(\b\d{2,4}\s*w\s*b\.?\s*\d+)',
    r'(\b\d{2,4}\s*w\s*\d+)',
    r'(Wolica)',
    r'(Obory)',
    r'(b\.?\s*2[2-4])'
]

SUBJECT_METADATA_PATTERNS = [
    r'(\b\d{1,2}[:.]\d{2}\b|\bod\s*\d{1,2}[:.]\d{2}!?\b|\bdo\s*\d{1,2}[:.]\d{2}!?\b)',
    r'(od\s*\d+\s*tygodnia(?:\s*zaj[eę]ć)?|1-\d+\s*tydzień(?:\s*zaj[eę]ć)?|2\s*poł\s*semestru|1\.\s*poł\s*semestru|\b1/2\b|\b2/2\b)',
    r'\bgr\.?\s*[\d,\s]+|\bgr\s*[\d,\s]+',
    r'\b\d{2,4}\s*/\s*Aula\s*\d*(?:\s*b\.?\s*\d+)?',
    r'sala\s*sekcyjna(?:\s*/\s*sala\s*ćwiczeniowa)?|sale\s*zakładu(?:\s*b\.?\s*\d+)?|(?:b\.?\s*\d+\s+)?sala\s*\d+[a-z]?(?:/\d+)?(?:\s*b\.?\s*\d+)?|aula\s*\d*|audytorium|aud\s*\d*|ambulatorium|ambul',
    r'\b\d{2,4}\s*w\s*b\.?\s*\d+|\b\d{2,4}\s*w\s*\d+|\bs\s*\d+\s*b\.?\s*\d+|\b06/ambul\b',
    r'b\.?\s*\d+',
    r'\bWolica\b|\bObory\b',
    r'\bćw\.?|\bćwiczenia|\bwykład\b|\bwyklad\b|\bfakultet\.?|\bw\.(?!\w)|\bw\b'
]


def extract_room_and_subject(text: str) -> Tuple[str, str, str]:
    """
    Wyodrębnia nazwę przedmiotu, salę oraz uwagi z tekstu komórki.
    """
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        return ("", "", "")

    full_text = " ".join(lines)

    # 1. Wykrywanie sali
    room = ""
    for rp in ROOM_PATTERNS:
        m = re.search(rp, full_text, re.IGNORECASE)
        if m:
            room = m.group(1).strip()
            break

    # 2. Wyodrębnienie i oczyszczenie nazwy przedmiotu
    first_line = lines[0]
    subject = first_line
    for pat in SUBJECT_METADATA_PATTERNS:
        subject = re.sub(pat, '', subject, flags=re.IGNORECASE)
    subject = re.sub(r'[*!,;:/()]+', ' ', subject)
    subject = ' '.join(subject.split()).strip()
    subject = re.sub(r'\s+[1-9]$', '', subject).strip()
    subject = re.sub(r'^[.\-\s]+|[.\-\s]+$', '', subject)

    if not subject:
        # Fallback do pierwszej linii bez dopisku formy
        subject = re.sub(r'\s*(?:wykład|wyklad|ćw\.?|cw\.?|ćwiczenia|fakultet).*$', '', first_line, flags=re.IGNORECASE).strip()
        if not subject:
            subject = first_line

    # 3. Uwagi / tygodnie
    notes = ""
    week_notes = re.findall(r'(?:od\s*\d+\s*tygodnia|1-\d+\s*tydzień|2\s*poł\s*semestru|1\.\s*poł\s*semestru|od\s*\d+:\d+!)', full_text, re.IGNORECASE)
    if week_notes:
        notes = ", ".join(week_notes)

    return (subject, room, notes)


def parse_sggw_rotations_excel(wb_or_sheet: Any) -> Dict[str, Dict[str, Dict[Any, Any]]]:
    """
    Parsuje arkusz rotacji klinicznych (semestr 11 VI rok WMW).
    Zwraca strukturę grup: 14 stażowych (gr. 1a-7b) oraz 7 ogólnych dziekańskich (gr. 1-7).
    """
    if isinstance(wb_or_sheet, str):
        wb = openpyxl.load_workbook(wb_or_sheet, data_only=True)
        sheet = wb.active
    elif hasattr(wb_or_sheet, "active"):
        sheet = wb_or_sheet.active
    else:
        sheet = wb_or_sheet

    clinics: Dict[int, str] = {}
    for col in range(3, min(20, sheet.max_column + 1)):
        v5 = sheet.cell(5, col).value
        v6 = sheet.cell(6, col).value
        parts = []
        if v5:
            parts.append(str(v5).strip())
        if v6 and str(v6).strip() not in parts:
            parts.append(str(v6).strip())
        name = " ".join(parts).strip()
        if name:
            clinics[col] = name

    subgroups = [f"gr. {i}{letter}" for i in range(1, 8) for letter in ('a', 'b')]
    main_groups = [f"gr. {i}" for i in range(1, 8)]
    all_groups = subgroups + main_groups

    result: Dict[str, Dict[str, Dict[Any, Any]]] = {
        g: {d: {} for d in ["PON", "WT", "ŚR", "CZW", "PT", "SOB", "ND"]} for g in all_groups
    }

    def _parse_range(text: str) -> Tuple[Optional[datetime.date], Optional[datetime.date]]:
        m = re.search(r'(\d{1,2})\.(\d{2})\s*[-–]\s*(\d{1,2})\.(\d{2})', text)
        if not m:
            return None, None
        d1, m1, d2, m2 = int(m.group(1)), int(m.group(2)), int(m.group(3)), int(m.group(4))
        y1 = 2026 if m1 >= 10 else 2027
        y2 = 2026 if m2 >= 10 else 2027
        return datetime(y1, m1, d1).date(), datetime(y2, m2, d2).date()

    weekday_map = {0: "PON", 1: "WT", 2: "ŚR", 3: "CZW", 4: "PT", 6: "ND"}

    # Wiersz 7: Wykłady ZZS dla wszystkich grup
    r7_c1 = str(sheet.cell(7, 1).value or "")
    d1_r7, d2_r7 = _parse_range(r7_c1)
    if d1_r7 and d2_r7:
        cur = d1_r7
        while cur <= d2_r7:
            if cur.weekday() in weekday_map:
                day_code = weekday_map[cur.weekday()]
                for g in all_groups:
                    slot_key = "wyklad_zzs"
                    result[g][day_code][slot_key] = {
                        "przedmiot": "Wykłady ZZS",
                        "raw_przedmiot": "Wykłady ZZS",
                        "prowadzacy": "",
                        "godziny": "08:00 - 14:30",
                        "sala": "CMT",
                        "height": 26,
                        "colspan": 1,
                        "data_start": d1_r7.isoformat(),
                        "tygodnie": 1,
                        "co_ile": 1,
                        "od_tyg": 1,
                        "polowa_sem": None,
                        "forma": "wyklad",
                        "uwagi": ""
                    }
            cur += timedelta(days=1)

    # Wiersze 8-25: Rotacje kliniczne
    for r in range(8, sheet.max_row + 1):
        c1_val = str(sheet.cell(r, 1).value or "")
        d1, d2 = _parse_range(c1_val)
        if not d1 or not d2:
            continue

        week_days = []
        cur = d1
        while cur <= d2:
            if cur.weekday() in weekday_map:
                week_days.append((cur, weekday_map[cur.weekday()]))
            cur += timedelta(days=1)

        for col, clinic_name in clinics.items():
            cell_val = sheet.cell(r, col).value
            if not cell_val:
                continue
            text = str(cell_val).strip()
            if text.lower() in ('x', '-', ''):
                continue

            found_subs = re.findall(r'(\d+)\s*([ab])', text, re.IGNORECASE)
            if not found_subs:
                continue

            notes_match = re.search(r'\(([^)]*(?:rozr[óo]d|CMT|staż)[^)]*)\)', text, re.IGNORECASE)
            notes = notes_match.group(1).strip() if notes_match else ""

            target_groups = set()
            for num, letter in found_subs:
                sub = f"gr. {num}{letter.lower()}"
                main = f"gr. {num}"
                target_groups.add(sub)
                target_groups.add(main)

            for d_date, day_code in week_days:
                slot_key = f"staz_{r}_{col}"
                subject_name = f"Staż - {clinic_name}"
                lesson_entry = {
                    "przedmiot": subject_name,
                    "raw_przedmiot": clinic_name,
                    "prowadzacy": "",
                    "godziny": "08:00 - 14:00",
                    "sala": clinic_name,
                    "height": 24,
                    "colspan": 1,
                    "data_start": d1.isoformat(),
                    "tygodnie": 1,
                    "co_ile": 1,
                    "od_tyg": 1,
                    "polowa_sem": None,
                    "forma": "staz",
                    "uwagi": notes
                }
                for g in target_groups:
                    result[g][day_code][slot_key] = lesson_entry

    return result


def parse_sggw_excel(wb_or_path: Union[str, openpyxl.Workbook]) -> Dict[str, Dict[str, Dict[Any, Any]]]:
    """
    Główna funkcja parsująca arkusz planu zajęć SGGW.
    Zwraca słownik grup i ich harmonogramów.
    """
    if isinstance(wb_or_path, str):
        wb = openpyxl.load_workbook(wb_or_path, data_only=True)
    else:
        wb = wb_or_path

    sheet = wb.active

    # Sprawdzenie czy arkusz to rotacje kliniczne (np. semestr 11)
    title_val = str(sheet.cell(1, 1).value or sheet.cell(1, 2).value or sheet.cell(1, 3).value or "")
    if "STAŻE" in title_val.upper() or "STAZE" in title_val.upper() or "ROTACJ" in title_val.upper():
        return parse_sggw_rotations_excel(sheet)

    # 1. Odczyt kolumn czasowych (szukamy w wierszach 1-3)
    time_columns: Dict[int, Tuple[str, str]] = {}
    time_row = 2
    for r_cand in [2, 1, 3]:
        for col in range(1, sheet.max_column + 1):
            v = sheet.cell(r_cand, col).value
            if v and re.search(r'\d{1,2}[:.]\d{2}\s*[-–]\s*\d{1,2}[:.]\d{2}', str(v)):
                time_row = r_cand
                break
        if any(re.search(r'\d{1,2}[:.]\d{2}\s*[-–]\s*\d{1,2}[:.]\d{2}', str(sheet.cell(time_row, c).value or "")) for c in range(1, min(10, sheet.max_column + 1))):
            break

    for col in range(1, sheet.max_column + 1):
        v = sheet.cell(time_row, col).value
        if v:
            v_str = str(v).strip()
            m = re.search(r'(\d{1,2}[:.]\d{2})\s*[-–]\s*(\d{1,2}[:.]\d{2})', v_str)
            if m:
                s = m.group(1).replace(".", ":")
                e = m.group(2).replace(".", ":")
                sh = f"{int(s.split(':')[0]):02d}:{s.split(':')[1]}"
                eh = f"{int(e.split(':')[0]):02d}:{e.split(':')[1]}"
                time_columns[col] = (sh, eh)

    if not time_columns:
        # Fallback dla arkuszy rotacji
        return parse_sggw_rotations_excel(sheet)

    # 2. Wykrycie liczby grup (z wiersza 1 lub domyślnie 7)
    title_val = str(sheet.cell(1, 1).value or sheet.cell(1, 2).value or sheet.cell(1, 3).value or "")
    m_groups = re.search(r'(\d+)\s*grup', title_val, re.IGNORECASE)
    total_groups_num = int(m_groups.group(1)) if m_groups else 7
    all_groups = [f"gr. {i}" for i in range(1, total_groups_num + 1)]

    # Przygotowanie struktury wyjściowej dla każdej grupy
    result: Dict[str, Dict[str, Dict[int, Any]]] = {
        g: {d: {} for d in ["PON", "WT", "ŚR", "CZW", "PT"]} for g in all_groups
    }

    # 3. Indeksowanie scalonych komórek
    merged_map: Dict[Tuple[int, int], Tuple[int, int, int, int]] = {}
    for rng in sheet.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged_map[(r, c)] = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)

    # 4. Przejście po wierszach i dniach tygodnia
    current_day = ""
    processed_top_lefts = set()
    start_c = min(time_columns.keys()) if time_columns else 2

    for r in range(time_row + 1, sheet.max_row + 1):
        # Sprawdzenie dnia w kolumnie 1 lub 2
        day_cell = sheet.cell(r, 1).value or sheet.cell(r, 2).value
        if day_cell:
            day_text = str(day_cell).strip().lower().split("\n")[0].rstrip(".:")
            for k, code in DAY_NAME_MAP.items():
                if day_text == k or day_text.startswith(k):
                    current_day = code
                    break

        if not current_day:
            continue

        for c in range(start_c, sheet.max_column + 1):
            cell_val = sheet.cell(r, c).value
            if not cell_val or not str(cell_val).strip():
                continue

            content = str(cell_val).strip()
            # Ignorowanie legendy / przypisów na dole
            if content.startswith("* Szczegółowy") or content.startswith("Plan zajęć"):
                continue

            # Sprawdzenie czy komórka jest częścią zakresu scalonego
            if (r, c) in merged_map:
                min_r, min_c, max_r, max_c = merged_map[(r, c)]
                if (min_r, min_c) in processed_top_lefts:
                    continue
                processed_top_lefts.add((min_r, min_c))
                start_col = min_c
                end_col = max_c
            else:
                start_col = c
                end_col = c

            # Wyznaczenie domyślnych godzin z kolumn
            def_start = time_columns.get(start_col, ("08:00", "09:00"))[0]
            def_end = time_columns.get(end_col, ("09:00", "10:00"))[1]

            actual_start, actual_end = parse_time_bounds(content, def_start, def_end)
            godziny_str = f"{actual_start} - {actual_end}"

            # Slot bazowy liczony jako minuty od 07:00 / 5
            try:
                sh_val, sm_val = map(int, actual_start.split(":"))
                eh_val, em_val = map(int, actual_end.split(":"))
                base_slot = ((sh_val - 7) * 60 + sm_val) // 5
                duration_min = (eh_val * 60 + em_val) - (sh_val * 60 + sm_val)
                height = duration_min // 5
            except Exception:
                base_slot = (start_col - 3) * 12 + 12
                height = (end_col - start_col + 1) * 12

            # Wyodrębnienie grup, przedmiotu i sali
            target_groups = extract_groups_from_text(content, all_groups)
            forma = classify_sggw_form(content)
            subject, room, notes = extract_room_and_subject(content)

            # Tygodnie i cykle
            weeks = 15
            co_ile = 1
            od_tyg = 1
            polowa_sem = None
            data_start = "2026-10-01"

            m_range_tyg = re.search(r'1\s*-\s*(\d+)\s*tydz', content, re.IGNORECASE)
            m_od_tyg = re.search(r'od\s*(\d+)\s*tygodnia', content, re.IGNORECASE)

            if m_range_tyg:
                weeks = int(m_range_tyg.group(1))
                data_start = "2026-10-01"
            elif m_od_tyg:
                od_tyg = int(m_od_tyg.group(1))
                weeks = max(1, 15 - od_tyg + 1)
                # Obliczenie daty startu: poniedziałek 1. tygodnia to 2026-10-05
                base_mon = datetime(2026, 10, 5)
                start_dt = base_mon + timedelta(days=(od_tyg - 1) * 7)
                data_start = start_dt.strftime("%Y-%m-%d")
            elif re.search(r'1[.\s]*poł(?:\.|owa)?\s*sem', content, re.IGNORECASE):
                weeks = 7
                polowa_sem = 1
                data_start = "2026-10-01"
            elif re.search(r'2[.\s]*poł(?:\.|owa)?\s*sem', content, re.IGNORECASE):
                weeks = 8
                polowa_sem = 2
                data_start = "2026-11-23"

            lesson_dict = {
                "przedmiot": subject,
                "raw_przedmiot": subject,
                "prowadzacy": "",
                "godziny": godziny_str,
                "sala": room or "brak sali",
                "height": height,
                "colspan": 1,
                "data_start": data_start,
                "tygodnie": weeks,
                "co_ile": co_ile,
                "od_tyg": od_tyg,
                "polowa_sem": polowa_sem,
                "forma": forma,
                "uwagi": notes
            }

            # Zapisanie do każdej zainteresowanej grupy
            for grp in target_groups:
                day_slots = result[grp][current_day]
                # Obsługa ewentualnych kolizji slotów (np. 2 zajęcia o tej samej godzinie)
                if base_slot not in day_slots:
                    day_slots[base_slot] = lesson_dict
                else:
                    existing = day_slots[base_slot]
                    if isinstance(existing, dict) and "przedmiot" in existing:
                        day_slots[base_slot] = {0: existing, 1: lesson_dict}
                    elif isinstance(existing, dict):
                        existing[len(existing)] = lesson_dict

    return result
